"""Reports, slides and scenarios generated from saved study artifacts."""
from __future__ import annotations

from pathlib import Path
import json
import shutil
from datetime import datetime, timezone

import pandas as pd

from market_forecast.dashboard.artifacts import ArtifactCatalog, load_catalog
from market_forecast.models import default_registry
from market_forecast.reports import audit_cross_market_report


def _markdown_table(frame: pd.DataFrame, columns: list[str]) -> str:
    selected = frame[columns].copy()
    for column in selected:
        if pd.api.types.is_float_dtype(selected[column]):
            selected[column] = selected[column].map(lambda value: "?" if pd.isna(value) else f"{value:.4g}")
    rows = ["| " + " | ".join(map(str, columns)) + " |", "| " + " | ".join(["---"] * len(columns)) + " |"]
    rows.extend("| " + " | ".join(str(value) for value in row) + " |" for row in selected.itertuples(index=False, name=None))
    return "\n".join(rows)


def _study_data(catalog: ArtifactCatalog) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    sources = []
    boundaries = []
    for run in catalog.runs:
        manifest = run.json("data_manifest.json")
        quality = run.json("quality_report.json")
        sources.append({
            "asset": run.asset_id, "run_id": run.run_id,
            "source": manifest["source_name"], "target": manifest["target_semantics"],
            "unit": manifest.get("unit", "USD"),
            "reuse": manifest["license_redistribution_status"],
            "rows": quality["row_count"],
        })
        boundaries.append({
            "asset": run.asset_id,
            "first validation": manifest["folds"][0]["test_start_date"][:10],
            "last validation": manifest["folds"][-1]["test_end_date"][:10],
            "holdout": f"{manifest['final_test_start_date'][:10]} to {manifest['final_test_end_date'][:10]}",
        })
    return pd.DataFrame(sources), pd.DataFrame(boundaries), catalog.table("recommendations.csv")


def build_scientific_report(catalog: ArtifactCatalog, output: Path) -> Path:
    audit = audit_cross_market_report(catalog.root)
    sources, boundaries, recommendations = _study_data(catalog)
    config = catalog.runs[0].json("config.json")
    failures = catalog.table("failures.csv")
    comparison = catalog.table("comparison.csv")
    quality_rows, diagnostic_rows, warning_rows = [], [], []
    for run in catalog.runs:
        quality = run.json("quality_report.json")
        eda = run.json("eda.json")
        quality_rows.append({
            "asset": run.asset_id, "first": quality["start"][:10], "last": quality["end"][:10],
            "duplicate keys": quality["duplicate_key_count"],
            "invalid OHLC": quality["inconsistent_ohlc_row_count"],
            "outlier flags": len(quality["close_return_outlier_dates"]),
        })
        diagnostic_rows.append({
            "asset": run.asset_id, "ADF p": eda["price"]["adf_pvalue"],
            "KPSS p": eda["price"]["kpss_pvalue"],
            "change std": eda["return_std"], "period": eda["seasonal_period"],
        })
        warning_rows.extend({"asset": run.asset_id, "warning": warning}
                            for warning in run.json("warnings.json"))
    columns = ["asset_id", "horizon", "model", "rmse_mean", "rmse_std", "final_test_rmse"]
    synthetic = (catalog.root / "SYNTHETIC_FIXTURE.txt").exists()
    title = "**SYNTHETIC OFFLINE FIXTURE ? SOFTWARE CHECK ONLY.**\n\n" if synthetic else ""
    text = f"""{title}# MarketCast Lab: {len(catalog.runs)} daily series

## Question and method

How does forecast performance vary by asset and observed-session horizon? Each series is evaluated separately with {config['folds']} expanding validation folds, a locked {config['final_test_size']}-observation holdout, and horizons {', '.join(map(str, config['horizons']))}. The forecast strategy is `{config['strategy']}`. Folds and holdout use each provider's observed dates; missing holidays and weekends are not manufactured.

The audit reproduced {audit['metric_rows_reproduced']} saved MAE/RMSE metric rows from fold predictions. Run IDs below link every result to its configuration, source hash, fold dates, warnings, and saved model parameters.

## Sources and targets

{_markdown_table(sources, ['asset', 'run_id', 'source', 'target', 'unit', 'reuse', 'rows'])}

Targets include adjusted prices, spot/reference values, and possibly percentage yields. Raw RMSE values across different units are not comparable. Source rights and provenance gaps are recorded in each run manifest.

## Cleaning, quality and exploratory analysis

{_markdown_table(pd.DataFrame(quality_rows), ['asset', 'first', 'last', 'duplicate keys', 'invalid OHLC', 'outlier flags'])}

Dates are sorted, duplicate asset/date keys rejected, and missing source prices dropped without adding synthetic market dates. Source OHLC fields remain null when only a point observation is published. Each run retains its raw source hash and a machine-readable quality report.

{_markdown_table(pd.DataFrame(diagnostic_rows), ['asset', 'ADF p', 'KPSS p', 'change std', 'period'])}

The EDA files include level and change diagnostics, rolling statistics, decomposition, and a stated transformation decision. Observation-count seasonal periods are exploratory proxies; holidays can shift their wall-clock span.

## Validation and holdout boundaries

{_markdown_table(boundaries, ['asset', 'first validation', 'last validation', 'holdout'])}

All preprocessing and model fitting take place within the applicable training history. Validation selects model parameters and recommendations; holdout values are displayed only after selection.

The configured model matrix is {', '.join(config['models'])}. MAE measures average absolute error, RMSE emphasizes larger errors, sMAPE scales by predicted and observed magnitude, and MASE compares with an in-sample naive difference. Runtime, residual diagnostics, and available statistical interval coverage are saved alongside the metrics.

## Selected models

{_markdown_table(recommendations, columns)}

A non-baseline model must improve both mean validation RMSE and mean plus one fold standard deviation by at least 5% over last value. Inspect `comparison.csv` and `fold_metrics.csv` for all models, fold variability, runtime and interval coverage. There were {len(failures)} recorded model failures. The run-level `warnings.json` files retain convergence and storage warnings.

{_markdown_table(comparison.head(20), ['asset_id', 'horizon', 'model', 'rmse_mean', 'rmse_std', 'final_test_rmse'])}

The table above is an excerpt of the full comparison file. Warnings recorded during the runs:

{_markdown_table(pd.DataFrame(warning_rows, columns=['asset', 'warning']), ['asset', 'warning']) if warning_rows else 'None recorded.'}

![Relative validation RMSE](relative_rmse.png)

![Fold variation by asset and horizon](fold_variability.png)

## Limits and reproducibility

The series have different calendars, target units, and publication mechanisms. A short historical window and one locked holdout limit ranking confidence. Recursive forecasts can compound error; available statistical intervals do not establish calibrated uncertainty. This is educational historical analysis, not a trading recommendation.

Regenerate runs with `python main.py four-markets`, audit with `python main.py audit-four-markets`, and rebuild deliverables with `python main.py build-deliverables`. The dashboard and report read the same `run_index.json` and saved predictions.
"""
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(text, encoding="utf-8")
    return output


def build_presentation(catalog: ArtifactCatalog, output: Path) -> Path:
    sources, boundaries, recommendations = _study_data(catalog)
    synthetic = (catalog.root / "SYNTHETIC_FIXTURE.txt").exists()
    lines = (["# SYNTHETIC OFFLINE FIXTURE", "Software check only.", "\n---\n"] if synthetic else []) + [
        "# MarketCast Lab", f"{len(catalog.runs)} daily series; saved backtests.",
        "\n---\n", "# Question and design", "How does forecast performance vary by asset and observed-session horizon?",
        "Validation folds select models; a locked holdout checks the selection.",
        "\n---\n", "# Sources and targets",
        _markdown_table(sources, ["asset", "source", "target", "unit"]),
        "\n---\n", "# Observation windows",
        _markdown_table(boundaries, ["asset", "first validation", "last validation", "holdout"]),
        "\n---\n", "# Selected models",
        _markdown_table(recommendations, ["asset_id", "horizon", "model", "rmse_mean", "final_test_rmse"]),
        "\n---\n", "# Interpretation and limits",
        "Compare models within an asset and target unit. Holdout results follow validation-only selection.",
        "Short historical windows and fold variation limit generalization. Educational analysis only.",
    ]
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return output


def build_powerpoint(catalog: ArtifactCatalog, output: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)

    def slide(title: str, body: str) -> None:
        item = deck.slides.add_slide(deck.slide_layouts[6])
        heading = item.shapes.add_textbox(Inches(.8), Inches(.5), Inches(11.8), Inches(.8))
        heading.text_frame.text = title
        heading.text_frame.paragraphs[0].font.size = Pt(30)
        box = item.shapes.add_textbox(Inches(.9), Inches(1.6), Inches(11.5), Inches(5.3))
        box.text_frame.word_wrap = True
        box.text_frame.text = body
        for paragraph in box.text_frame.paragraphs:
            paragraph.font.size = Pt(19)

    label = "Synthetic fixture: software check only." if (catalog.root / "SYNTHETIC_FIXTURE.txt").exists() else "Historical evaluation only."
    slide("MarketCast Lab", f"{len(catalog.runs)} daily series\n{label}")
    slide("Question and design", "How does forecast performance vary by asset and observed-session horizon?\nValidation selects models; the locked holdout checks them.")
    for start in range(0, len(catalog.runs), 6):
        slide("Sources and targets", "\n".join(
            f"{run.asset_id}: {run.json('data_manifest.json')['target_semantics']}"
            for run in catalog.runs[start:start + 6]))
    recommendations = catalog.table("recommendations.csv")
    for run in catalog.runs:
        selected = recommendations[recommendations.asset_id == run.asset_id]
        slide(run.asset_id, "\n".join(
            f"Horizon {int(row.horizon)}: {row.model}; validation RMSE {row.rmse_mean:.4g}; holdout RMSE {row.final_test_rmse:.4g}"
            for row in selected.itertuples()))
    slide("Interpretation", "Compare models within each target unit. Inspect fold variation, source manifests and holdout reversals.\nHistorical analysis is not investment advice.")
    output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(output)
    return output


def build_deliverables(index_root: str | Path = "assets/phase3", output_root: str | Path = "assets/phase4") -> Path:
    catalog = load_catalog(index_root)
    output = Path(output_root)
    if not output.is_absolute():
        output = catalog.root.parents[1] / output
    output.mkdir(parents=True, exist_ok=True)
    build_scientific_report(catalog, output / "scientific_report.md")
    build_presentation(catalog, output / "presentation.md")
    try:
        build_powerpoint(catalog, output / "presentation.pptx")
    except ImportError:
        pass  # Markdown slides remain available without the optional presentation extra.
    for name in ("relative_rmse.png", "fold_variability.png"):
        shutil.copy2(catalog.root / name, output / name)
    return output


def build_future_scenarios(index_root: str | Path = "assets/phase3", output_root: str | Path = "assets/phase4") -> Path:
    """Explicitly refit selected models to the saved study endpoint for scenario display."""
    catalog = load_catalog(index_root)
    output = Path(output_root)
    if not output.is_absolute():
        output = catalog.root.parents[1] / output
    output.mkdir(parents=True, exist_ok=True)
    recommendations = catalog.table("recommendations.csv")
    registry = default_registry()
    rows = []
    training_rows = {}
    for run in catalog.runs:
        config = run.json("config.json")
        manifest = run.json("data_manifest.json")
        frame = run.source_data()
        if config["max_train_rows"]:
            frame = frame.iloc[-int(config["max_train_rows"]):]
        target = manifest["target_column"]
        values = frame[target].to_numpy(dtype=float)
        training_rows[run.asset_id] = len(values)
        for selected in recommendations[recommendations.asset_id == run.asset_id].itertuples():
            params = dict(manifest["selected_model_parameters"].get(selected.model, {}))
            if selected.model in {"ridge", "lasso", "random_forest", "xgboost", "rnn", "lstm", "gru"}:
                params.setdefault("seed", config["seed"])
            model = registry.create(selected.model, params).fit(values)
            prediction = model.predict(int(selected.horizon))
            rows.extend({"asset_id": run.asset_id, "horizon": int(selected.horizon),
                         "step": step, "prediction": float(value), "model": selected.model,
                         "origin_timestamp": frame.timestamp.iloc[-1].isoformat(),
                         "source_run_id": run.run_id, "target": target}
                        for step, value in enumerate(prediction, start=1))
    path = output / "future_forecasts.csv"
    pd.DataFrame(rows).to_csv(path, index=False)
    (output / "future_manifest.json").write_text(json.dumps({
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_index": str(catalog.root / "run_index.json"),
        "method": "explicit recursive refit through each saved study cutoff",
        "training_rows_by_asset": training_rows,
        "calendar_dates": "omitted; steps count observations without invented future holidays",
        "purpose": "speculative educational scenario, separate from historical backtest",
    }, indent=2), encoding="utf-8")
    return path
